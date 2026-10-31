# -*- coding: utf-8 -*-
"""
Module Front Matter cho Báo cáo Đồ án NoSQL - Trường Đại học Công Thương TP. Hồ Chí Minh (HUIT)
Bao gồm:
1. Trang bìa chính (Hard Cover - có khung viền hoa văn kép, in trên giấy bìa)
2. Trang bìa phụ (Soft Cover - in giấy trắng thường)
3. Lời cảm ơn (Acknowledgements)
4. Phiếu nhận xét của Giảng viên hướng dẫn
5. Phiếu nhận xét của Giảng viên phản biện
6. Mục lục chi tiết
7. Danh mục từ viết tắt
8. Danh mục bảng biểu
9. Danh mục hình ảnh

Định dạng tuân thủ 100% quy định Khoa CNTT - HUIT:
- Font: Times New Roman
- Số trang: Không đánh số trang bìa chính và bìa phụ.
- Đánh số La Mã (i, ii, iii...) từ Lời cảm ơn đến hết Danh mục hình ảnh.
- Bảng: Không màu nền, chữ đen 100%, viền đen mảnh đơn giản.
"""

import os
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BLACK = RGBColor(0, 0, 0)

def set_section_cover_borders(section):
    """Đặt khung viền mỹ thuật kép cho trang bìa chính"""
    sectPr = section._sectPr
    for b in sectPr.findall(qn('w:pgBorders')):
        sectPr.remove(b)
    pgBorders = OxmlElement('w:pgBorders')
    pgBorders.set(qn('w:offsetFrom'), 'page')
    for b_name in ('top', 'left', 'bottom', 'right'):
        border = OxmlElement(f'w:{b_name}')
        border.set(qn('w:val'), 'double') # Viền kép nghệ thuật
        border.set(qn('w:sz'), '12')       # 1.5 pt
        border.set(qn('w:space'), '24')    # Khoảng cách lề mép trang
        border.set(qn('w:color'), '000000')# Màu đen
        pgBorders.append(border)
    sectPr.append(pgBorders)

def clear_section_borders(section):
    """Xóa bỏ khung viền trang cho các section nội dung"""
    sectPr = section._sectPr
    for b in sectPr.findall(qn('w:pgBorders')):
        sectPr.remove(b)

def set_section_page_number_type(section, fmt="decimal", start=1):
    """Cấu hình định dạng đánh số trang (lowerRoman hoặc decimal) và số bắt đầu"""
    sectPr = section._sectPr
    for p in sectPr.findall(qn('w:pgNumType')):
        sectPr.remove(p)
    if fmt:
        pgNumType = OxmlElement('w:pgNumType')
        pgNumType.set(qn('w:fmt'), fmt)
        if start is not None:
            pgNumType.set(qn('w:start'), str(start))
        sectPr.append(pgNumType)

def add_page_number_field(run):
    """Chèn trường số trang PAGE tự động vào Run"""
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def add_cover_page(doc, helpers, is_main_cover=True):
    """
    Tạo trang bìa:
    is_main_cover = True: Trang bìa chính (có khung viền mỹ thuật kép)
    is_main_cover = False: Trang bìa phụ (không khung viền ngoài)
    """
    section = doc.sections[-1]
    
    # Cấu hình lề trang bìa: Trái 3.5cm, Phải 2.5cm, Trên 2.5cm, Dưới 2.5cm
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    
    # Không đánh số trang cho bìa
    section.footer.is_linked_to_previous = False
    section.footer.paragraphs[0].text = ''
    set_section_page_number_type(section, fmt=None, start=None)
    
    # Khung viền trang cho bìa
    if is_main_cover:
        set_section_cover_borders(section)
    else:
        clear_section_borders(section)
        
    # Cơ quan chủ quản và Tên trường
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header.paragraph_format.line_spacing = 1.2
    p_header.paragraph_format.space_before = Pt(0)
    p_header.paragraph_format.space_after = Pt(2)
    p_header.paragraph_format.first_line_indent = Cm(0)
    
    r_hdr1 = p_header.add_run("BỘ CÔNG THƯƠNG\n")
    r_hdr1.font.name = 'Times New Roman'
    r_hdr1.font.size = Pt(13)
    r_hdr1.bold = True
    r_hdr1.font.color.rgb = BLACK
    
    r_hdr2 = p_header.add_run("TRƯỜNG ĐẠI HỌC CÔNG THƯƠNG TP. HỒ CHÍ MINH\n")
    r_hdr2.font.name = 'Times New Roman'
    r_hdr2.font.size = Pt(14)
    r_hdr2.bold = True
    r_hdr2.font.color.rgb = BLACK
    
    r_hdr3 = p_header.add_run("KHOA CÔNG NGHỆ THÔNG TIN - BỘ MÔN CƠ SỞ DỮ LIỆU\n")
    r_hdr3.font.name = 'Times New Roman'
    r_hdr3.font.size = Pt(13)
    r_hdr3.bold = True
    r_hdr3.font.color.rgb = BLACK
    
    # Dấu gạch phân cách
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(14)
    p_line.paragraph_format.first_line_indent = Cm(0)
    r_l = p_line.add_run("-------------------o0o-------------------")
    r_l.font.name = 'Times New Roman'
    r_l.font.size = Pt(12)
    r_l.font.color.rgb = BLACK
    
    # Logo HUIT
    assets_dir = helpers.get('assets_dir', '')
    logo_path = os.path.join(assets_dir, 'image1.jpeg')
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(4)
        p_logo.paragraph_format.space_after = Pt(12)
        p_logo.paragraph_format.first_line_indent = Cm(0)
        run_logo = p_logo.add_run()
        run_logo.add_picture(logo_path, width=Cm(3.5))
        
    # Tiêu đề báo cáo
    p_report = doc.add_paragraph()
    p_report.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_report.paragraph_format.line_spacing = 1.3
    p_report.paragraph_format.space_before = Pt(8)
    p_report.paragraph_format.space_after = Pt(6)
    p_report.paragraph_format.first_line_indent = Cm(0)
    
    r_rep1 = p_report.add_run("BÁO CÁO ĐỒ ÁN HỌC PHẦN\n")
    r_rep1.font.name = 'Times New Roman'
    r_rep1.font.size = Pt(16)
    r_rep1.bold = True
    r_rep1.font.color.rgb = BLACK
    
    r_rep2 = p_report.add_run("HỌC PHẦN: DỮ LIỆU NOSQL (MÃ HP: 01201073)\n\n")
    r_rep2.font.name = 'Times New Roman'
    r_rep2.font.size = Pt(14)
    r_rep2.bold = True
    r_rep2.font.color.rgb = BLACK
    
    r_rep3 = p_report.add_run("ĐỀ TÀI:\n")
    r_rep3.font.name = 'Times New Roman'
    r_rep3.font.size = Pt(15)
    r_rep3.bold = True
    r_rep3.font.color.rgb = BLACK
    
    r_title = p_report.add_run("NỀN TẢNG E-LEARNING & KHẢO SÁT NĂNG LỰC TIẾNG ANH THỰC CHIẾN TÍCH HỢP ĐA MÔ HÌNH NOSQL\n(EDUPRO ACADEMY)\n")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(18)
    r_title.bold = True
    r_title.font.color.rgb = BLACK
    
    # Bảng thông tin GVHD và Sinh viên thực hiện
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(12)
    p_space.paragraph_format.space_after = Pt(4)
    p_space.paragraph_format.first_line_indent = Cm(0)
    
    tbl_info = doc.add_table(rows=5, cols=2)
    tbl_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w = [Cm(5.0), Cm(10.0)]
    
    info_data = [
        ("Giảng viên hướng dẫn:", "Bộ môn Hệ thống Thông tin"),
        ("Sinh viên thực hiện:", "1. Nguyễn Trọng Nguyên  - MSSV: 2001230589 (Trưởng nhóm)\n2. Phạm Nguyễn Khôi Nguyên - MSSV: 2001230578\n3. Võ Huỳnh Gia Huy    - MSSV: 2001230303"),
        ("Lớp học phần:", "14DHTH"),
        ("Khóa học:", "14 (2023 - 2027)"),
        ("Hệ đào tạo:", "Chính quy")
    ]
    
    for r_idx, (label, val) in enumerate(info_data):
        row = tbl_info.rows[r_idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        
        helpers['set_cell_margins'](cell_lbl, top=50, bottom=50, left=60, right=60)
        helpers['set_cell_margins'](cell_val, top=50, bottom=50, left=60, right=60)
        
        # Bảng thông tin bìa không viền
        p_l = cell_lbl.paragraphs[0]
        p_l.paragraph_format.first_line_indent = Cm(0)
        p_l.paragraph_format.line_spacing = 1.3
        p_l.paragraph_format.space_before = Pt(2)
        p_l.paragraph_format.space_after = Pt(2)
        run_l = p_l.add_run(label)
        run_l.font.name = 'Times New Roman'
        run_l.font.size = Pt(13)
        run_l.bold = True
        run_l.font.color.rgb = BLACK
        
        p_v = cell_val.paragraphs[0]
        p_v.paragraph_format.first_line_indent = Cm(0)
        p_v.paragraph_format.line_spacing = 1.3
        p_v.paragraph_format.space_before = Pt(2)
        p_v.paragraph_format.space_after = Pt(2)
        run_v = p_v.add_run(val)
        run_v.font.name = 'Times New Roman'
        run_v.font.size = Pt(13)
        run_v.font.color.rgb = BLACK
        if r_idx == 0:
            run_v.bold = True
            
        row.cells[0].width = col_w[0]
        row.cells[1].width = col_w[1]
        
    # Địa điểm và thời gian
    p_loc = doc.add_paragraph()
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_loc.paragraph_format.space_before = Pt(24)
    p_loc.paragraph_format.space_after = Pt(0)
    p_loc.paragraph_format.first_line_indent = Cm(0)
    r_loc = p_loc.add_run("TP. HỒ CHÍ MINH, THÁNG 10 NĂM 2026")
    r_loc.font.name = 'Times New Roman'
    r_loc.font.size = Pt(13)
    r_loc.bold = True
    r_loc.font.color.rgb = BLACK

def add_acknowledgements(doc, helpers):
    """Tạo trang LỜI CẢM ƠN trang trọng, súc tích chuẩn học thuật HUIT"""
    helpers['add_heading_chapter'](doc, "LỜI CẢM ƠN")
    
    helpers['add_paragraph_huit'](doc, 
        "Sau thời gian học tập, nghiên cứu và nỗ lực triển khai thực tế, đồ án học phần Dữ liệu NoSQL với đề tài \"Nền tảng E-Learning & Khảo sát năng lực tiếng Anh thực chiến tích hợp đa mô hình NoSQL (EduPro Academy)\" đã được nhóm sinh viên hoàn thiện. Để đạt được kết quả này, bên cạnh sự nỗ lực cố gắng của các thành viên trong nhóm, chúng em đã may mắn nhận được sự quan tâm giúp đỡ quý báu, sự tận tâm giảng dạy từ Quý Thầy Cô, cùng sự ủng hộ to lớn từ Gia đình và Bạn bè."
    )
    
    helpers['add_paragraph_huit'](doc, 
        "Trước hết, nhóm chúng em xin gửi lời cảm ơn chân thành, sâu sắc và lòng kính trọng nhất tới Bộ môn Hệ thống Thông tin – giảng viên trực tiếp giảng dạy và tận tâm hướng dẫn học phần Dữ liệu NoSQL (Mã HP: 01201073). Trong suốt quá trình thực hiện đồ án, Thầy/Cô đã luôn dành nhiều thời gian quý báu để định hướng phương pháp tiếp cận đề tài, truyền đạt những nền tảng tri thức chuyên sâu về cơ sở dữ liệu phi quan hệ và các kiến trúc phân tán hiện đại. Những lời góp ý khoa học, thực tiễn và sự chỉ dẫn ân cần của Giảng viên chính là kim chỉ nam giúp nhóm vượt qua các rào cản kỹ thuật phức tạp, hoàn thiện sản phẩm đúng tiến độ và nâng cao chất lượng học thuật của toàn bộ công trình."
    )
    
    helpers['add_paragraph_huit'](doc, 
        "Nhóm tác giả cũng xin trân trọng gửi lời cảm ơn đến Ban Giám hiệu Trường Đại học Công Thương TP. Hồ Chí Minh (HUIT), Ban Chủ nhiệm Khoa Công nghệ Thông tin cùng toàn thể Quý Thầy, Cô giáo thuộc Bộ môn Cơ sở Dữ liệu. Quý Thầy, Cô đã luôn tâm huyết truyền dạy cho chúng em hành trang kiến thức công nghệ thông tin vững chắc, xây dựng môi trường học thuật chuyên nghiệp và trang bị hệ thống phòng thực hành máy tính hiện đại, tạo nền tảng tối ưu cho việc nghiên cứu và thực nghiệm đề tài."
    )
    
    helpers['add_paragraph_huit'](doc, 
        "Chúng con xin bày tỏ lòng biết ơn vô hạn đến Cha Mẹ và Gia đình – điểm tựa tinh thần vững chắc nhất, đã luôn thương yêu, thấu hiểu, động viên và tạo mọi điều kiện tốt nhất để chúng con an tâm học tập, nghiên cứu và theo đuổi đam mê. Đồng thời, nhóm cũng xin gửi lời cảm ơn chân thành tới các bạn sinh viên tập thể lớp 14DHTH đã luôn đồng hành, chia sẻ kinh nghiệm học tập và tích cực hỗ trợ nhóm trong các đợt kiểm thử thực tế hệ thống phần mềm."
    )
    
    helpers['add_paragraph_huit'](doc, 
        "Mặc dù các thành viên trong nhóm đã dành trọn tâm huyết và làm việc với tinh thần nghiêm túc cao nhất, song do giới hạn về mặt thời gian và kinh nghiệm thực tiễn, đồ án chắc chắn khó tránh khỏi những thiếu sót nhất định. Nhóm sinh viên thực hiện rất mong nhận được những lời chỉ bảo, đánh giá chân thành và những ý kiến đóng góp quý báu từ Quý Thầy, Cô trong Hội đồng chấm điểm để đề tài tiếp tục được hoàn thiện và phát triển hơn nữa trong tương lai."
    )
    
    # Chữ ký nhóm tác giả
    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign.paragraph_format.space_before = Pt(16)
    p_sign.paragraph_format.space_after = Pt(2)
    p_sign.paragraph_format.first_line_indent = Cm(0)
    p_sign.paragraph_format.line_spacing = 1.3
    
    r_date = p_sign.add_run("TP. Hồ Chí Minh, ngày 30 tháng 10 năm 2026\n")
    r_date.font.name = 'Times New Roman'
    r_date.font.size = Pt(13)
    r_date.italic = True
    r_date.font.color.rgb = BLACK
    
    r_team = p_sign.add_run("Nhóm sinh viên thực hiện (Nhóm 18)\n\n\n\n")
    r_team.font.name = 'Times New Roman'
    r_team.font.size = Pt(13)
    r_team.bold = True
    r_team.font.color.rgb = BLACK
    
    r_names = p_sign.add_run("Nguyễn Trọng Nguyên\nPhạm Nguyễn Khôi Nguyên\nVõ Huỳnh Gia Huy")
    r_names.font.name = 'Times New Roman'
    r_names.font.size = Pt(13)
    r_names.bold = True
    r_names.font.color.rgb = BLACK

def add_advisor_review(doc, helpers):
    """Tạo Phiếu nhận xét của Giảng viên hướng dẫn"""
    helpers['add_heading_chapter'](doc, "NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN")
    
    helpers['add_paragraph_huit'](doc, 
        "Cán bộ hướng dẫn: Bộ môn Hệ thống Thông tin\n"
        "Đơn vị công tác: Khoa Công nghệ Thông tin - Trường Đại học Công Thương TP. Hồ Chí Minh\n"
        "Học phần: Dữ liệu NoSQL (Mã HP: 01201073)\n"
        "Tên đề tài: Nền tảng E-Learning & Khảo sát năng lực tiếng Anh thực chiến tích hợp đa mô hình NoSQL (EduPro Academy)\n"
        "Nhóm sinh viên thực hiện: Nhóm 18 - Lớp 14DHTH\n"
        "1. Nguyễn Trọng Nguyên (MSSV: 2001230589 - Trưởng nhóm)\n"
        "2. Phạm Nguyễn Khôi Nguyên (MSSV: 2001230578)\n"
        "3. Võ Huỳnh Gia Huy (MSSV: 2001230303)",
        "1. Thông tin chung: "
    )
    
    helpers['add_heading_level1'](doc, "2. NỘI DUNG NHẬN XÉT VÀ ĐÁNH GIÁ")
    
    table_headers = ["STT", "Nội dung đánh giá", "Mức độ đánh giá / Điểm tối đa", "Điểm đánh giá"]
    table_data = [
        ["1", "Tinh thần, thái độ và mức độ chuyên cần của nhóm trong các buổi làm việc", "Tốt / 1.0 điểm", "........ / 1.0"],
        ["2", "Mức độ hoàn thành mục tiêu nghiên cứu và khối lượng công việc theo tiến độ", "Tốt / 2.0 điểm", "........ / 2.0"],
        ["3", "Khả năng phân tích yêu cầu (SRS) và thiết kế kiến trúc đa mô hình NoSQL (SDS)", "Xuất sắc / 2.5 điểm", "........ / 2.5"],
        ["4", "Chất lượng sản phẩm ứng dụng, tính ổn định và mức độ thực chiến của 4 hệ NoSQL", "Xuất sắc / 2.5 điểm", "........ / 2.5"],
        ["5", "Hình thức trình bày báo cáo, tính khoa học và độ tin cậy của kết quả thực nghiệm", "Chuẩn mực / 2.0 điểm", "........ / 2.0"],
        ["TỔNG", "TỔNG ĐIỂM ĐÁNH GIÁ CỦA GIẢNG VIÊN HƯỚNG DẪN", "Thang điểm 10", "........ / 10.0"]
    ]
    helpers['create_huit_table'](doc, "Bảng đánh giá kết quả thực hiện đồ án của Giảng viên hướng dẫn", table_headers, table_data, [1.2, 8.0, 3.8, 2.5])
    
    helpers['add_heading_level1'](doc, "3. Ý KIẾN NHẬN XÉT CHI TIẾT CỦA GIẢNG VIÊN HƯỚNG DẪN")
    helpers['add_paragraph_huit'](doc, "....................................................................................................................................................................................")
    helpers['add_paragraph_huit'](doc, "....................................................................................................................................................................................")
    helpers['add_paragraph_huit'](doc, "....................................................................................................................................................................................")
    helpers['add_paragraph_huit'](doc, "....................................................................................................................................................................................")
    
    helpers['add_paragraph_huit'](doc, 
        "Đề tài có / không đáp ứng đầy đủ yêu cầu của học phần Dữ liệu NoSQL và đủ điều kiện / không đủ điều kiện để được bảo vệ trước Hội đồng chấm điểm đồ án.",
        "Kết luận: "
    )
    
    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign.paragraph_format.space_before = Pt(14)
    p_sign.paragraph_format.space_after = Pt(2)
    p_sign.paragraph_format.first_line_indent = Cm(0)
    p_sign.paragraph_format.line_spacing = 1.3
    
    r_date = p_sign.add_run("TP. Hồ Chí Minh, ngày ...... tháng ...... năm 2026\n")
    r_date.font.name = 'Times New Roman'
    r_date.font.size = Pt(13)
    r_date.italic = True
    r_date.font.color.rgb = BLACK
    
    r_gv = p_sign.add_run("Giảng viên hướng dẫn\n(Ký và ghi rõ họ tên)\n\n\n\n\n")
    r_gv.font.name = 'Times New Roman'
    r_gv.font.size = Pt(13)
    r_gv.bold = True
    r_gv.font.color.rgb = BLACK
    
    r_name = p_sign.add_run("Bộ môn Hệ thống Thông tin")
    r_name.font.name = 'Times New Roman'
    r_name.font.size = Pt(13)
    r_name.bold = True
    r_name.font.color.rgb = BLACK

def add_reviewer_review(doc, helpers):
    """Tạo Phiếu nhận xét của Giảng viên phản biện"""
    helpers['add_heading_chapter'](doc, "NHẬN XÉT CỦA GIẢNG VIÊN PHẢN BIỆN")
    
    helpers['add_paragraph_huit'](doc, 
        "Cán bộ chấm phản biện: ...........................................................................................................................\n"
        "Đơn vị công tác: Khoa Công nghệ Thông tin - Trường Đại học Công Thương TP. Hồ Chí Minh\n"
        "Tên đề tài: Nền tảng E-Learning & Khảo sát năng lực tiếng Anh thực chiến tích hợp đa mô hình NoSQL (EduPro Academy)\n"
        "Nhóm sinh viên thực hiện: Nhóm 18 - Lớp 14DHTH\n"
        "1. Nguyễn Trọng Nguyên (MSSV: 2001230589 - Trưởng nhóm)\n"
        "2. Phạm Nguyễn Khôi Nguyên (MSSV: 2001230578)\n"
        "3. Võ Huỳnh Gia Huy (MSSV: 2001230303)",
        "1. Thông tin chung: "
    )
    
    helpers['add_heading_level1'](doc, "2. NỘI DUNG NHẬN XÉT VÀ ĐÁNH GIÁ CỦA GIẢNG VIÊN PHẢN BIỆN")
    
    table_headers = ["STT", "Tiêu chuẩn đánh giá phản biện", "Mức độ đánh giá / Điểm tối đa", "Điểm đánh giá"]
    table_data = [
        ["1", "Tính cấp thiết, tính khoa học và giá trị ứng dụng thực tiễn của đề tài", "Phù hợp / 1.5 điểm", "........ / 1.5"],
        ["2", "Tính hợp lý trong việc lựa chọn và phân chia vai trò của 4 hệ CSDL NoSQL", "Chính xác / 2.5 điểm", "........ / 2.5"],
        ["3", "Mức độ đáp ứng của 20 câu truy vấn (10 cơ bản + 10 nâng cao) trên GUI Tool", "Đầy đủ / 2.0 điểm", "........ / 2.0"],
        ["4", "Hiệu năng hệ thống, kiến trúc kết nối ứng dụng và kết quả kiểm thử tự động", "Xuất sắc / 2.0 điểm", "........ / 2.0"],
        ["5", "Chất lượng quyển báo cáo, cấu trúc trình bày và khả năng trả lời phản biện", "Tốt / 2.0 điểm", "........ / 2.0"],
        ["TỔNG", "TỔNG ĐIỂM ĐÁNH GIÁ CỦA GIẢNG VIÊN PHẢN BIỆN", "Thang điểm 10", "........ / 10.0"]
    ]
    helpers['create_huit_table'](doc, "Bảng đánh giá kết quả đồ án của Giảng viên phản biện", table_headers, table_data, [1.2, 8.0, 3.8, 2.5])
    
    helpers['add_heading_level1'](doc, "3. CÂU HỎI PHẢN BIỆN DÀNH CHO NHÓM SINH VIÊN KHI BẢO VỆ")
    helpers['add_paragraph_huit'](doc, "Câu 1: ...........................................................................................................................................................................")
    helpers['add_paragraph_huit'](doc, "Câu 2: ...........................................................................................................................................................................")
    helpers['add_paragraph_huit'](doc, "Câu 3: ...........................................................................................................................................................................")
    
    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign.paragraph_format.space_before = Pt(14)
    p_sign.paragraph_format.space_after = Pt(2)
    p_sign.paragraph_format.first_line_indent = Cm(0)
    p_sign.paragraph_format.line_spacing = 1.3
    
    r_date = p_sign.add_run("TP. Hồ Chí Minh, ngày ...... tháng ...... năm 2026\n")
    r_date.font.name = 'Times New Roman'
    r_date.font.size = Pt(13)
    r_date.italic = True
    r_date.font.color.rgb = BLACK
    
    r_gv = p_sign.add_run("Giảng viên phản biện\n(Ký và ghi rõ họ tên)\n\n\n\n\n")
    r_gv.font.name = 'Times New Roman'
    r_gv.font.size = Pt(13)
    r_gv.bold = True
    r_gv.font.color.rgb = BLACK
    
    r_name = p_sign.add_run("......................................................")
    r_name.font.name = 'Times New Roman'
    r_name.font.size = Pt(13)
    r_name.bold = True
    r_name.font.color.rgb = BLACK

def add_table_of_contents(doc, helpers):
    """Tạo bảng MỤC LỤC chi tiết phân cấp chuẩn HUIT"""
    helpers['add_heading_chapter'](doc, "MỤC LỤC")
    
    toc_headers = ["Đề mục", "Tên đề mục nội dung", "Trang"]
    toc_data = [
        ["", "LỜI CẢM ƠN", "i"],
        ["", "NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN", "ii"],
        ["", "NHẬN XÉT CỦA GIẢNG VIÊN PHẢN BIỆN", "iii"],
        ["", "MỤC LỤC", "iv"],
        ["", "DANH MỤC TỪ VIẾT TẮT", "vi"],
        ["", "DANH MỤC BẢNG BIỂU", "vii"],
        ["", "DANH MỤC HÌNH ẢNH", "viii"],
        ["", "PHẦN MỞ ĐẦU", "1"],
        ["1", "Giới thiệu bối cảnh và lý do chọn đề tài", "1"],
        ["2", "Mục tiêu đề tài", "2"],
        ["3", "Đối tượng và phạm vi đề tài", "3"],
        ["4", "Phương pháp nghiên cứu", "4"],
        ["5", "Bố cục đồ án", "5"],
        ["CHƯƠNG 1", "TỔNG QUAN ĐỀ TÀI VÀ KẾ HOẠCH DỰ ÁN", "6"],
        ["1.1", "Bối cảnh và lý do chọn đề tài", "6"],
        ["1.2", "Mục tiêu đề tài", "7"],
        ["1.3", "Phạm vi nghiên cứu và triển khai", "8"],
        ["1.4", "Phương pháp thực hiện", "8"],
        ["1.5", "Công nghệ và công cụ sử dụng", "9"],
        ["1.6", "Kế hoạch dự án và phân công nhiệm vụ", "10"],
        ["CHƯƠNG 2", "TỔNG QUAN NOSQL VÀ TÌM HIỂU TOOL", "13"],
        ["2.1", "Tổng quan cơ sở dữ liệu NoSQL", "13"],
        ["2.2", "Phân tán dữ liệu trong NoSQL", "15"],
        ["2.3", "Phân loại các mô hình NoSQL", "17"],
        ["2.4", "Các hệ quản trị cơ sở dữ liệu đề tài sử dụng", "18"],
        ["2.5", "Tìm hiểu và so sánh công cụ giao diện GUI Tool", "22"],
        ["CHƯƠNG 3", "ĐẶC TẢ YÊU CẦU PHẦN MỀM (SRS)", "24"],
        ["3.1", "Giới thiệu chuẩn IEEE 830 / ISO 29148", "24"],
        ["3.2", "Mô tả tổng quan hệ thống", "25"],
        ["3.3", "Yêu cầu chức năng (Functional Requirements - 15 FRs)", "26"],
        ["3.3.1", "Mô hình hóa nghiệp vụ (BPMN Swimlane)", "28"],
        ["3.3.2", "Mô hình hóa chức năng (Use Case, Activity, Sequence, State Machine, DFD)", "29"],
        ["3.4", "Yêu cầu phi chức năng (Non-Functional Requirements - 6 NFRs)", "33"],
        ["3.5", "Yêu cầu dữ liệu đặc thù NoSQL (3V & Conceptual Model)", "34"],
        ["3.6", "Mẫu truy cập dữ liệu (Access Patterns - AP-01 đến AP-16)", "36"],
        ["3.7", "Yêu cầu giao diện (UI, API Endpoints, Môi trường hạ tầng)", "38"],
        ["3.8", "Ma trận truy vết yêu cầu (FR - AP - CSDL - TC)", "40"],
        ["CHƯƠNG 4", "ĐẶC TẢ THIẾT KẾ PHẦN MỀM (SDS)", "42"],
        ["4.1", "Kiến trúc tổng thể và mô hình triển khai Docker Container", "42"],
        ["4.2", "Quyết định lựa chọn loại CSDL (Ma trận quyết định NoSQL)", "43"],
        ["4.3", "Thiết kế CSDL Key-Value (Redis In-Memory)", "44"],
        ["4.4", "Thiết kế CSDL Document (MongoDB)", "45"],
        ["4.5", "Thiết kế CSDL Column-family (Apache Cassandra)", "47"],
        ["4.6", "Thiết kế CSDL Graph (Neo4j)", "49"],
        ["4.7", "Đồng bộ dữ liệu giữa các CSDL (Polyglot Persistence Architecture)", "50"],
        ["4.8", "Phân tán, sao chép và định lý CAP/PACELC", "51"],
        ["4.9", "Bảo mật, sao lưu và giám sát hệ thống", "52"],
        ["4.10", "Thiết kế Module và danh mục API Backend", "53"],
        ["4.11", "Ma trận ánh xạ Access Pattern (AP) sang thiết kế CSDL", "55"],
        ["4.12", "Mô hình hóa cấu trúc (Class Diagram & Entity Mapping)", "56"],
        ["4.13", "Các kỹ thuật tối ưu nâng cao", "57"],
        ["CHƯƠNG 5", "XÂY DỰNG CSDL VÀ TRUY VẤN TRÊN GUI TOOL", "59"],
        ["5.1", "Môi trường triển khai cụm CSDL và GUI Tool", "59"],
        ["5.2", "Khởi tạo Database, Schema Validation, Keyspaces và Constraints", "60"],
        ["5.3", "Nạp dữ liệu thực tế cho ứng dụng (Data Seeding)", "62"],
        ["5.4", "Triển khai 10 câu truy vấn cơ bản trên 4 hệ NoSQL", "63"],
        ["5.5", "Triển khai 10 câu truy vấn nâng cao trên 4 hệ NoSQL", "66"],
        ["5.6", "Sao lưu và phục hồi dữ liệu (Backup & Restore Evidence)", "70"],
        ["CHƯƠNG 6", "KẾT NỐI ỨNG DỤNG VÀ DEMO", "72"],
        ["6.1", "Kiến trúc phân tầng ứng dụng (Layered Architecture)", "72"],
        ["6.2", "Cấu hình Driver và kết nối 4 CSDL NoSQL", "73"],
        ["6.3", "Demo 6 nhóm chức năng nghiệp vụ thực tế kèm minh chứng", "75"],
        ["6.4", "Hiện thực hóa các kỹ thuật nâng cao trong thực tế", "81"],
        ["6.5", "Phần mở rộng hệ thống", "82"],
        ["CHƯƠNG 7", "KIỂM THỬ VÀ ĐÁNH GIÁ HIỆU NĂNG", "84"],
        ["7.1", "Kiểm thử chức năng (Bộ kiểm thử tự động 22 Test Cases)", "84"],
        ["7.2", "Đánh giá hiệu năng k6 và 5 thí nghiệm so sánh thực nghiệm", "86"],
        ["7.3", "Kiểm thử khả năng chịu lỗi (Chaos Engineering & High Availability)", "89"],
        ["7.4", "Thảo luận kết quả và phân tích đánh đổi (Trade-offs)", "90"],
        ["CHƯƠNG 8", "KẾT QUẢ ĐẠT ĐƯỢC VÀ TỔNG KẾT HOẠT ĐỘNG NHÓM", "91"],
        ["8.1", "Đánh giá mức độ hoàn thành mục tiêu đồ án", "91"],
        ["8.2", "Đánh giá mức độ đáp ứng các yêu cầu phi chức năng (NFR)", "92"],
        ["8.3", "Tổng kết hoạt động nhóm và tỷ lệ đóng góp của từng thành viên", "93"],
        ["8.4", "Bài học kinh nghiệm rút ra từ đề tài", "94"],
        ["8.5", "Hướng phát triển và mở rộng trong tương lai", "95"],
        ["", "TÀI LIỆU THAM KHẢO", "96"]
    ]
    helpers['create_huit_table'](doc, "Mục lục đồ án tốt nghiệp học phần NoSQL", toc_headers, toc_data, [2.5, 11.5, 1.5])

def add_abbreviations(doc, helpers):
    """Tạo DANH MỤC TỪ VIẾT TẮT chuẩn chỉnh"""
    helpers['add_heading_chapter'](doc, "DANH MỤC TỪ VIẾT TẮT")
    
    abbr_headers = ["Từ viết tắt", "Thuật ngữ tiếng Anh", "Ý nghĩa / Định nghĩa tiếng Việt"]
    abbr_data = [
        ["ACID", "Atomicity, Consistency, Isolation, Durability", "Bốn thuộc tính giao dịch truyền thống trong RDBMS (Nguyên tử, Nhất quán, Cô lập, Bền vững)"],
        ["AP", "Access Pattern", "Mẫu truy cập dữ liệu (Mẫu câu hỏi ứng dụng đặt ra cho cơ sở dữ liệu)"],
        ["API", "Application Programming Interface", "Giao diện lập trình ứng dụng"],
        ["AOF", "Append Only File", "Cơ chế ghi nhật ký thay đổi tuần tự trên đĩa của hệ quản trị Redis"],
        ["BASE", "Basically Available, Soft-state, Eventual consistency", "Mô hình đảm bảo tính sẵn sàng cơ bản và nhất quán cuối cùng của NoSQL"],
        ["BPMN", "Business Process Model and Notation", "Ký hiệu và mô hình hóa quy trình nghiệp vụ chuẩn quốc tế"],
        ["CAP", "Consistency, Availability, Partition tolerance", "Định lý đánh đổi giữa Nhất quán, Sẵn sàng và Chịu lỗi phân vùng mạng"],
        ["CMS", "Content Management System", "Hệ thống quản trị nội dung dành cho Quản trị viên và Giảng viên"],
        ["CQL", "Cassandra Query Language", "Ngôn ngữ truy vấn dạng bảng của hệ quản trị Apache Cassandra"],
        ["CRUD", "Create, Read, Update, Delete", "Bốn thao tác cơ bản trên dữ liệu: Tạo mới, Đọc, Cập nhật và Xóa"],
        ["DFD", "Data Flow Diagram", "Sơ đồ luồng dữ liệu mô tả sự chuyển dịch thông tin trong hệ thống"],
        ["ESR", "Equality, Sort, Range", "Quy tắc vàng thiết kế chỉ mục kết hợp tối ưu trên MongoDB"],
        ["FR", "Functional Requirement", "Yêu cầu chức năng của phần mềm theo chuẩn IEEE 830"],
        ["GUI", "Graphical User Interface", "Giao diện người dùng đồ họa"],
        ["HUIT", "Ho Chi Minh City University of Industry and Trade", "Trường Đại học Công Thương Thành phố Hồ Chí Minh"],
        ["JWT", "JSON Web Token", "Chuẩn mã hóa thẻ xác thực phiên đăng nhập phi trạng thái"],
        ["LMS", "Learning Management System", "Hệ thống quản lý học tập trực tuyến"],
        ["LPG", "Labeled Property Graph", "Mô hình đồ thị thuộc tính có gán nhãn của hệ quản trị Neo4j"],
        ["NFR", "Non-Functional Requirement", "Yêu cầu phi chức năng (Hiệu năng, Sẵn sàng, Bảo mật...)"],
        ["NoSQL", "Not Only SQL", "Thuật ngữ chỉ các hệ quản trị cơ sở dữ liệu phi quan hệ hiện đại"],
        ["PACELC", "Partition-Availability/Consistency-Else-Latency/Consistency", "Định lý mở rộng của CAP giải thích sự đánh đổi giữa Độ trễ và Nhất quán khi mạng bình thường"],
        ["RDBMS", "Relational Database Management System", "Hệ quản trị cơ sở dữ liệu quan hệ truyền thống (MySQL, SQL Server...)"],
        ["RDB", "Redis Database Backup", "Cơ chế sao lưu chụp ảnh nhanh bộ nhớ RAM của Redis định kỳ xuống đĩa"],
        ["SDS", "Software Design Specification", "Tài liệu đặc tả thiết kế phần mềm theo chuẩn IEEE 1016"],
        ["SRS", "Software Requirements Specification", "Tài liệu đặc tả yêu cầu phần mềm theo chuẩn IEEE 830 / ISO 29148"],
        ["TC", "Test Case", "Trường hợp kiểm thử phần mềm"],
        ["TTL", "Time To Live", "Thời gian tồn tại tự động hết hạn của dữ liệu trong Redis và Cassandra"],
        ["TWCS", "Time Window Compaction Strategy", "Chiến lược nén ghép nối tệp dữ liệu theo cửa sổ thời gian trong Cassandra"],
        ["UI/UX", "User Interface / User Experience", "Giao diện người dùng và Trải nghiệm người dùng"],
        ["XP", "Experience Point", "Điểm kinh nghiệm tích lũy qua hoạt động học tập dùng đổi quà Gamification"]
    ]
    helpers['create_huit_table'](doc, "Danh mục các thuật ngữ và từ viết tắt trong đồ án", abbr_headers, abbr_data, [2.2, 5.8, 7.5])

def add_list_of_tables(doc, helpers):
    """Tạo DANH MỤC BẢNG BIỂU chuẩn HUIT"""
    helpers['add_heading_chapter'](doc, "DANH MỤC BẢNG BIỂU")
    
    tbl_headers = ["Ký hiệu bảng", "Tên bảng biểu", "Trang"]
    tbl_data = [
        ["Bảng 1.1", "Kế hoạch thực hiện đề tài theo mô hình Agile Scrum (12 tuần)", "10"],
        ["Bảng 1.2", "Bảng phân công nhiệm vụ và trách nhiệm từng thành viên Nhóm 18", "11"],
        ["Bảng 1.3", "Danh mục công nghệ và công cụ phát triển đồ án", "12"],
        ["Bảng 2.1", "Ma trận so sánh đặc tính kỹ thuật giữa RDBMS và 4 mô hình NoSQL", "14"],
        ["Bảng 2.2", "So sánh công cụ quản trị GUI Tool cho 4 hệ CSDL NoSQL", "23"],
        ["Bảng 3.1", "Danh sách các tác nhân hệ thống và đặc điểm người dùng", "25"],
        ["Bảng 3.2", "Danh mục 15 yêu cầu chức năng (Functional Requirements - FR)", "26"],
        ["Bảng 3.3", "Danh mục 6 yêu cầu phi chức năng (Non-Functional Requirements - NFR)", "33"],
        ["Bảng 3.4", "Đặc tính dữ liệu theo mô hình 3V (Volume, Velocity, Variety) và TTL", "34"],
        ["Bảng 3.5", "Từ điển dữ liệu các thực thể cốt lõi trong hệ thống", "35"],
        ["Bảng 3.6", "Danh mục 16 mẫu truy cập dữ liệu (Access Patterns - AP-01 đến AP-16)", "36"],
        ["Bảng 3.7", "Danh mục các giao diện lập trình ứng dụng cốt lõi (REST API Endpoints)", "38"],
        ["Bảng 3.8", "Ma trận truy vết yêu cầu (Requirement Traceability Matrix: FR - AP - CSDL - TC)", "40"],
        ["Bảng 4.1", "Ma trận quyết định lựa chọn CSDL NoSQL cho từng nhóm dữ liệu", "43"],
        ["Bảng 4.2", "Thiết kế cấu trúc Key-Value trên Redis", "44"],
        ["Bảng 4.3", "Danh mục Collection và quyết định nhúng/tham chiếu trên MongoDB", "45"],
        ["Bảng 4.4", "Thiết kế hệ thống chỉ mục (Index) trên MongoDB theo quy tắc ESR", "46"],
        ["Bảng 4.5", "Thiết kế các bảng Column-family trên Apache Cassandra", "47"],
        ["Bảng 4.6", "Danh mục các nhãn Node và thuộc tính trên Neo4j", "49"],
        ["Bảng 4.7", "Danh mục các loại quan hệ (Relationships) trên Neo4j", "49"],
        ["Bảng 4.8", "Nguồn sự thật (Source of Truth) và cơ chế đồng bộ dữ liệu đa CSDL", "50"],
        ["Bảng 4.9", "Phân loại dữ liệu theo Định lý CAP và PACELC", "51"],
        ["Bảng 4.10", "Chính sách bảo mật, phân quyền và lịch sao lưu dữ liệu NoSQL", "52"],
        ["Bảng 4.11", "Danh mục chi tiết 15 API Backend của hệ thống", "54"],
        ["Bảng 4.12", "Ma trận ánh xạ Mẫu truy cập dữ liệu (AP) sang thiết kế vật lý CSDL", "55"],
        ["Bảng 4.13", "Ánh xạ cấu trúc dữ liệu NoSQL sang lớp đối tượng Backend (Entity Mapping)", "56"],
        ["Bảng 5.1", "Môi trường triển khai cụm CSDL NoSQL và công cụ GUI Tool", "59"],
        ["Bảng 5.2", "Danh mục 10 câu truy vấn cơ bản trên 4 hệ quản trị NoSQL", "63"],
        ["Bảng 5.3", "Danh mục 10 câu truy vấn nâng cao trên 4 hệ quản trị NoSQL", "66"],
        ["Bảng 6.1", "Ma trận liên kết chức năng ứng dụng với 4 hệ quản trị NoSQL", "75"],
        ["Bảng 7.1", "Tổng hợp kết quả kiểm thử tự động 22 Test Cases chức năng (TC-01 đến TC-22)", "84"],
        ["Bảng 7.2", "Kết quả kiểm thử tải k6 cho các kịch bản đọc/ghi dồn dập", "86"],
        ["Bảng 7.3", "Bảng tổng hợp kết quả 5 thí nghiệm so sánh hiệu năng thực nghiệm", "87"],
        ["Bảng 8.1", "Đánh giá mức độ hoàn thành các chức năng hệ thống so với mục tiêu", "91"],
        ["Bảng 8.2", "Đánh giá mức độ đáp ứng các yêu cầu phi chức năng (NFR)", "92"],
        ["Bảng 8.3", "Bảng tổng kết hoạt động và mức độ đóng góp cá nhân của từng thành viên", "93"]
    ]
    helpers['create_huit_table'](doc, "Danh mục các bảng biểu trong đồ án", tbl_headers, tbl_data, [2.5, 11.5, 1.5])

def add_list_of_figures(doc, helpers):
    """Tạo DANH MỤC HÌNH ẢNH chuẩn HUIT"""
    helpers['add_heading_chapter'](doc, "DANH MỤC HÌNH ẢNH")
    
    fig_headers = ["Ký hiệu hình", "Tên hình ảnh minh họa", "Trang"]
    fig_data = [
        ["Hình 5.1", "Minh chứng sao lưu dữ liệu toàn diện (Full Backup) trên hệ thống Polyglot NoSQL", "71"],
        ["Hình 6.1", "Giao diện Trang chủ và Quản lý danh mục khóa học tiếng Anh (MongoDB Document Store)", "76"],
        ["Hình 6.2", "Giao diện Chi tiết khóa học và Danh sách bài học video (MongoDB)", "77"],
        ["Hình 6.3", "Giao diện Sổ tay từ vựng Flashcard cá nhân hóa (MongoDB)", "78"],
        ["Hình 6.4", "Giao diện Bảng xếp hạng học tập thời gian thực Leaderboard (Redis Sorted Set)", "79"],
        ["Hình 6.5", "Giao diện Phòng thi trắc nghiệm trực tuyến (MongoDB Core & Redis Session Cache)", "80"],
        ["Hình 6.6", "Hệ thống chống gian lận Anti-cheat 3/3 cảnh báo (Redis Event Tracking)", "80"],
        ["Hình 6.7", "Giao diện Cửa hàng đổi quà Gamification bằng điểm thưởng XP (Redis & MongoDB)", "81"],
        ["Hình 6.8", "Giao diện Quản trị viên CMS phê duyệt nội dung và ngân hàng câu hỏi (MongoDB)", "81"]
    ]
    helpers['create_huit_table'](doc, "Danh mục các hình ảnh minh chứng trong đồ án", fig_headers, fig_data, [2.5, 11.5, 1.5])
