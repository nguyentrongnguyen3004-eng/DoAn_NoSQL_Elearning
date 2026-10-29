# -*- coding: utf-8 -*-
"""
Script điền đầy đủ và chuẩn hóa BẢNG BÁO CÁO MINH CHỨNG VÀ TỰ ĐÁNH GIÁ ĐỒ ÁN NoSQL
Tuân thủ 100% thể thức Khoa CNTT - HUIT:
- Font Times New Roman, chữ màu đen
- Điền đầy đủ thông tin 3 sinh viên Nhóm 18
- Đầy đủ 7 mục hồ sơ nộp
- Đầy đủ phân công nhiệm vụ và tóm tắt kiến trúc 4 CSDL NoSQL
"""

import sys, os
sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BLACK = RGBColor(0, 0, 0)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border_black(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        border = OxmlElement(f'w:{edge}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tcBorders.append(border)
    tcPr.append(tcBorders)

def format_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=12):
    cell.text = text
    set_cell_margins(cell)
    set_cell_border_black(cell)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Cm(0)
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(font_size)
        r.bold = bold
        r.font.color.rgb = BLACK

def generate_minh_chung_report():
    template_path = r'D:\TRONGNGUYEN_HK7\NoSQL\ĐỒ ÁN\TÀI LIỆU ZALO\BaoCaoMinhChung_DoAn_DuLieuNoSQL.docx'
    doc = docx.Document(template_path)
    
    # Chuẩn hóa cấu hình lề trang
    for s in doc.sections:
        s.left_margin = Cm(3.5)
        s.right_margin = Cm(2.5)
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)

    # Chuẩn hóa font các đoạn văn tiêu đề
    for p in doc.paragraphs:
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.first_line_indent = Cm(0)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.color.rgb = BLACK

    # =========================================================
    # TABLE 0: THÔNG TIN CHUNG
    # =========================================================
    t0 = doc.tables[0]
    info_map = {
        0: ("Tên đồ án", "Nền tảng E-Learning & Khảo sát năng lực tiếng Anh thực chiến tích hợp đa mô hình NoSQL (EduPro Academy)"),
        1: ("Mã nhóm", "Nhóm 18 - Lớp 14DHTH"),
        2: ("Danh sách nhóm", "1. Nguyễn Trọng Nguyên - MSSV: 2001230589 (Trưởng nhóm)\n2. Phạm Nguyễn Khôi Nguyên - MSSV: 2001230578\n3. Võ Huỳnh Gia Huy - MSSV: 2001230303"),
        3: ("Liên kết mã nguồn", "https://github.com/nguyentrongnguyen3004-eng/DoAn_NoSQL_Elearning.git"),
        4: ("Các liên kết trong báo cáo", "• Video Demo: Đính kèm thư mục 07_Video_Demo và link Google Drive\n• Giao diện Web LMS: http://localhost:5050\n• Bản báo cáo hoàn chỉnh: BAO_CAO_NOSQL_EDUPRO_ACADEMY_FINAL.docx")
    }
    for r_idx, (lbl, val) in info_map.items():
        row = t0.rows[r_idx]
        format_cell_text(row.cells[0], lbl, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=12)
        format_cell_text(row.cells[1], val, bold=(r_idx in [0, 1]), align=WD_ALIGN_PARAGRAPH.LEFT, font_size=12)

    # =========================================================
    # TABLE 1: DANH MỤC HỒ SƠ PHẢI NỘP (7 MỤC)
    # =========================================================
    t1 = doc.tables[1]
    # Header
    format_cell_text(t1.rows[0].cells[0], "STT", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)
    format_cell_text(t1.rows[0].cells[1], "Hồ sơ / Sản phẩm", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)
    format_cell_text(t1.rows[0].cells[2], "Tên tệp hoặc đường dẫn", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)
    format_cell_text(t1.rows[0].cells[3], "Đã nộp", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)
    format_cell_text(t1.rows[0].cells[4], "Ghi chú", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)

    items = [
        ("1", "Source code đầy đủ", "01_Source_Code_EduPro/", "Đã nộp", "Node.js 20+, Express RESTful API, Frontend Vanilla ES6 module hóa tại public/js/"),
        ("2", "Script cho từng hệ CSDL NoSQL", "02_Scripts_CSDL_NoSQL/", "Đã nộp", "Đầy đủ 4 hệ: mongo_seed.js, redis_demo.js, cassandra_schema.cql, neo4j_graph.cypher"),
        ("3", "Tài liệu đặc tả yêu cầu phần mềm", "03_Tai_Lieu_Do_An/03_Dac_Ta_Yeu_Cau_SRS/\nNhom_18_SRS_Elearning.docx", "Đã nộp", "Đặc tả SRS chuẩn IEEE 830 / ISO 29148 với 15 FRs, 16 APs, 6 NFRs, BPMN, DFD, Sequence"),
        ("4", "Tài liệu đặc tả thiết kế phần mềm", "03_Tai_Lieu_Do_An/04_Dac_Ta_Thiet_Ke_Phan_Mem/\nNhom_18_GiaiDoan2_Database_Model.docx", "Đã nộp", "Đặc tả SDS chuẩn IEEE 1016, thiết kế chi tiết 4 CSDL NoSQL, PACELC, 15 API, Class Diagram"),
        ("5", "Tài liệu/kịch bản kiểm thử", "03_Tai_Lieu_Do_An/05_Tai_Lieu_Kiem_Thu/\nautomated_tests.js", "Đã nộp", "22/22 kịch bản Test Case tự động đạt 100% Passed, k6 load testing, 5 thí nghiệm so sánh"),
        ("6", "Báo cáo kết quả thực hiện đồ án", "TÀI LIỆU FINAL/\nBAO_CAO_NOSQL_EDUPRO_ACADEMY_FINAL.docx", "Đã nộp", "Báo cáo toàn văn chuẩn HUIT (Bìa, Lời cảm ơn, Nhận xét, Mục lục, Mở đầu, Chương 1-8, TLTK)"),
        ("7", "Video demo các tính năng của sản phẩm", "07_Video_Demo/LINK_VIDEO_DEMO.txt", "Đã nộp", "Thời lượng 5-7 phút theo đúng cấu trúc rubric chấm điểm trực tiếp của giảng viên")
    ]
    for idx, (stt, prod, path, status, note) in enumerate(items):
        row = t1.rows[idx + 1]
        format_cell_text(row.cells[0], stt, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)
        format_cell_text(row.cells[1], prod, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=12)
        format_cell_text(row.cells[2], path, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=11)
        format_cell_text(row.cells[3], status, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)
        format_cell_text(row.cells[4], note, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=11)

    # =========================================================
    # TABLE 2: THÔNG TIN NHÓM VÀ PHÂN CÔNG (3 THÀNH VIÊN)
    # =========================================================
    t2 = doc.tables[2]
    # Header
    format_cell_text(t2.rows[0].cells[0], "STT", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)
    format_cell_text(t2.rows[0].cells[1], "Họ và tên - MSSV", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)
    format_cell_text(t2.rows[0].cells[2], "Công việc thực hiện", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)
    format_cell_text(t2.rows[0].cells[3], "Tỷ lệ đóng góp dự kiến (%)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)

    members = [
        ("1", "Nguyễn Trọng Nguyên\nMSSV: 2001230589\n(Trưởng nhóm)", 
         "Nhóm trưởng, thiết kế kiến trúc tổng thể Polyglot NoSQL, cấu hình MongoDB Core Store & Redis Cache, xây dựng phân hệ Quản lý đào tạo LMS (Khóa học, Bài học video, Từ vựng Flashcard, Sổ tay thông minh), xây dựng bộ kiểm thử tự động 22 Test Cases, biên soạn Chương 1, 2, 4, 7 và tổng hợp báo cáo.",
         "40%"),
        ("2", "Phạm Nguyễn Khôi Nguyên\nMSSV: 2001230578", 
         "Thiết kế Keyspace và bảng Column-family trên Apache Cassandra, xây dựng hệ thống theo dõi tiến độ học tập chuỗi thời gian (Time-series Learning Events, Heartbeat 5s), triển khai phòng thi trắc nghiệm trực tuyến thời gian thực và thuật toán chống gian lận Anti-cheat 3/3 cảnh báo, đo lường kiểm thử tải k6, biên soạn Chương 3, 5.",
         "30%"),
        ("3", "Võ Huỳnh Gia Huy\nMSSV: 2001230303", 
         "Thiết kế mô hình đồ thị thuộc tính (Labeled Property Graph) trên Neo4j, xây dựng thuật toán gợi ý Bạn Cùng Tiến (Study Buddies) và Đồ thị bài học phụ thuộc (Prerequisite Path), xây dựng hệ sinh thái Gamification đổi quà bằng điểm thưởng XP trên Redis & MongoDB, thực hiện sao lưu/phục hồi dữ liệu, biên soạn Chương 6, 8.",
         "30%")
    ]
    for idx, (stt, name, task, pct) in enumerate(members):
        row = t2.rows[idx + 1]
        format_cell_text(row.cells[0], stt, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)
        format_cell_text(row.cells[1], name, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=12)
        format_cell_text(row.cells[2], task, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=11)
        format_cell_text(row.cells[3], pct, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)

    # =========================================================
    # TABLE 3: TÓM TẮT ĐỀ TÀI VÀ KIẾN TRÚC GIẢI PHÁP
    # =========================================================
    t3 = doc.tables[3]
    summary_map = {
        0: ("Bài toán cần giải quyết", 
            "Giải quyết bài toán lưu trữ và xử lý đa chiều trong hệ thống E-Learning đào tạo tiếng Anh thực chiến có lượng truy cập lớn và dữ liệu biến động cao: loại bỏ hoàn toàn các điểm nghẽn của RDBMS truyền thống về khóa dòng khi nộp bài dồn dập, quá tải I/O đĩa từ khi ghi log hành vi chuỗi thời gian (Write-Heavy), cấu trúc bài giảng linh hoạt đa định dạng và chi phí bùng nổ khi JOIN đệ quy duyệt đồ thị lộ trình học tập."),
        1: ("Đối tượng sử dụng", 
            "Bao gồm 3 nhóm tác nhân chính:\n• Học viên (Student): Tham gia học tập, rèn luyện từ vựng, làm bài thi khảo sát trực tuyến, đổi quà Gamification bằng XP.\n• Giảng viên (Teacher): Quản lý khóa học, biên tập ngân hàng câu hỏi, tạo bài kiểm tra, theo dõi tỷ lệ giữ chân video học viên.\n• Quản trị viên (Admin): Giám sát hệ thống, kiểm duyệt nội dung, quản lý phân quyền, cấu hình an ninh và bảo mật."),
        2: ("Chức năng chính", 
            "• Khảo sát năng lực Placement Test & Thi chứng chỉ (IELTS, TOEIC) thời gian thực\n• Hệ thống chống gian lận thi cử thông minh Anti-cheat 3/3 lần cảnh báo\n• Sổ tay từ vựng Flashcard cá nhân hóa & Phòng thực hành AI Writing Lab\n• Bảng xếp hạng thi đua tuần thời gian thực (Real-time Leaderboard)\n• Gợi ý lộ trình bài học phụ thuộc (Learning Path) & Mạng lưới bạn học cùng tiến (Study Buddies)\n• Ghi nhận tiến độ học tập chuỗi thời gian phân tích tỷ lệ giữ chân video (Time-series Event Logging)\n• Cửa hàng Gamification đổi vật phẩm học tập bằng điểm thưởng XP"),
        3: ("Kiến trúc, Công nghệ sử dụng", 
            "• Frontend: HTML5, CSS3/Tailwind CSS, JavaScript ES6 (Kiến trúc Module hóa tại public/js/)\n• Backend: Node.js 20+, Express.js RESTful API, Docker & Docker Compose\n• Testing: Node Test Runner tích hợp 22 Test Cases, k6 load testing, Chaos Engineering"),
        4: ("RDBMS sử dụng (nếu có)", 
            "Không sử dụng (Hệ thống kiến trúc hoàn toàn 100% trên nền tảng Polyglot NoSQL đa mô hình)"),
        5: ("Các hệ NoSQL sử dụng", 
            "1. MongoDB (Document Store): Kho lưu trữ bền vững (Users, Courses, Lessons, Exams, Questions, Flashcards, Results, Certificates). Áp dụng quy tắc chỉ mục ESR và Schema Validation.\n2. Redis (Key-Value Store): In-Memory Cache siêu tốc, Blacklist JWT, Session thi trực tuyến với TTL, Real-time Leaderboard bằng Sorted Set, Sliding Window Rate Limiter.\n3. Apache Cassandra (Column-family Store): Lưu trữ chuỗi thời gian ghi tải cực cao (learning_events, violation_logs) với chiến lược nén TWCS và tự động xóa bằng TTL.\n4. Neo4j (Graph Database): Biểu diễn Đồ thị kỹ năng (Skill Graph), Mạng lưới lộ trình bài học (REQUIRES), và Thuật toán gợi ý bạn cùng tiến (ENROLLED_IN)."),
        6: ("Các sơ đồ Database:", 
            "• Sơ đồ kiến trúc tổng thể Polyglot Persistence NoSQL Architecture\n• Sơ đồ Document Data Model & Index ESR trên MongoDB\n• Sơ đồ Key-Value & In-Memory Data Structures trên Redis\n• Sơ đồ Chebotko Column-family trên Apache Cassandra CQL\n• Sơ đồ Labeled Property Graph (Nodes & Relationships) trên Neo4j Cypher\n(Được trình bày và đặc tả chi tiết tại Chương 4 trong quyển báo cáo toàn văn)")
    }
    for r_idx, (lbl, val) in summary_map.items():
        row = t3.rows[r_idx]
        format_cell_text(row.cells[0], lbl, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=12)
        format_cell_text(row.cells[1], val, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=11)

    out_paths = [
        r'D:\TRONGNGUYEN_HK7\NoSQL\ĐỒ ÁN\TÀI LIỆU FINAL\BaoCaoMinhChung_DoAn_DuLieuNoSQL_HoanChinh.docx',
        r'D:\TRONGNGUYEN_HK7\NoSQL\ĐỒ ÁN\03_Tai_Lieu_Do_An\06_Bao_Cao_Ket_Qua_Do_An\BaoCaoMinhChung_DoAn_DuLieuNoSQL_HoanChinh.docx'
    ]
    for p in out_paths:
        doc.save(p)
        print(f"✅ Đã lưu Báo cáo minh chứng hoàn chỉnh tại: {p}")

if __name__ == '__main__':
    generate_minh_chung_report()
