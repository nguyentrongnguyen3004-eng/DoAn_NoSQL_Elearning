import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = pptx.Presentation('03_Tai_Lieu_Do_An/BAO_CAO_THUYET_TRINH_NOSQL_EDUPRO_ACADEMY.pptx')
slide = prs.slides[19] # Slide 20

# Find and remove Rounded Rectangle 9
to_remove = None
for sh in slide.shapes:
    if sh.name == 'Rounded Rectangle 9':
        to_remove = sh
        break

if to_remove:
    sp = to_remove._element
    sp.getparent().remove(sp)
    print("Removed old text box 9")

# Add title box for table
title_box = slide.shapes.add_textbox(Inches(0.80), Inches(3.70), Inches(11.73), Inches(0.40))
p = title_box.text_frame.paragraphs[0]
p.text = "BẢNG ĐỐI SÁNH HIỆU NĂNG THỰC NGHIỆM (K6 & AUTOCANNON BENCHMARK TRÊN 10.000 VIRTUAL USERS)"
p.font.name = "Segoe UI"
p.font.size = Pt(11.5)
p.font.bold = True
p.font.color.rgb = RGBColor(79, 70, 229)

# Add native PowerPoint table
rows = 6
cols = 6
left = Inches(0.80)
top = Inches(4.15)
width = Inches(11.73)
height = Inches(2.50)

table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
table = table_shape.table

# Set column widths
table.columns[0].width = Inches(3.60)
table.columns[1].width = Inches(2.10)
table.columns[2].width = Inches(1.60)
table.columns[3].width = Inches(1.20)
table.columns[4].width = Inches(1.20)
table.columns[5].width = Inches(2.03)

headers = ["Mẫu truy cập (Access Pattern)", "Công nghệ NoSQL", "Throughput (req/s)", "Độ trễ p50", "Độ trễ p95", "Đánh giá NFR"]
data = [
    ["AP-01: Đọc danh mục khóa học (Cache Hit)", "Redis Cache-Aside", "4,250 req/s", "2.1 ms", "3.8 ms", "✓ Đạt xuất sắc (NFR <50ms)"],
    ["AP-01: Đọc danh mục từ đĩa (No Cache)", "MongoDB Collection", "520 req/s", "24.5 ms", "38.2 ms", "✓ Đạt tiêu chuẩn"],
    ["AP-06: Ghi nhật ký video Telemetry", "Apache Cassandra", "2,480 writes/s", "3.9 ms", "6.4 ms", "✓ Đạt xuất sắc (Không khóa)"],
    ["AP-07: Đua top bảng xếp hạng tuần (ZSET)", "Redis Sorted Set", "3,890 req/s", "1.8 ms", "3.2 ms", "✓ Đạt xuất sắc (O(log N))"],
    ["AP-09: Tìm lộ trình học bài tiên quyết", "Neo4j ShortestPath", "1,120 req/s", "3.4 ms", "5.8 ms", "✓ Đạt xuất sắc (O(1) Adjacency)"]
]

for col_idx, h_text in enumerate(headers):
    cell = table.cell(0, col_idx)
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(79, 70, 229) # Indigo
    p = cell.text_frame.paragraphs[0]
    p.text = h_text
    p.font.name = "Segoe UI"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER if col_idx >= 2 else PP_ALIGN.LEFT

for r_idx, row_data in enumerate(data):
    for c_idx, val in enumerate(row_data):
        cell = table.cell(r_idx + 1, c_idx)
        cell.fill.solid()
        if r_idx % 2 == 0:
            cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
        else:
            cell.fill.fore_color.rgb = RGBColor(241, 245, 249) # Slate 100
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.name = "Segoe UI"
        p.font.size = Pt(9.0)
        p.font.color.rgb = RGBColor(15, 23, 42)
        if c_idx >= 2 and c_idx <= 4:
            p.alignment = PP_ALIGN.CENTER
            p.font.bold = (c_idx == 2)
        elif c_idx == 5:
            p.alignment = PP_ALIGN.CENTER
            p.font.bold = True
            p.font.color.rgb = RGBColor(5, 150, 105) # Emerald green

prs.save('03_Tai_Lieu_Do_An/BAO_CAO_THUYET_TRINH_NOSQL_EDUPRO_ACADEMY.pptx')
prs.save('TÀI LIỆU FINAL/BAO_CAO_THUYET_TRINH_NOSQL_EDUPRO_ACADEMY.pptx')
print("Successfully converted Slide 20 benchmark table to native PowerPoint table!")
